# Конвертер Google-таблиці «Журнал» у резервну копію Каси (формат /api/import)
# Використання: pip install openpyxl; python3 tools/import_xlsx.py таблиця.xlsx kasa-import.json [дата-звіту-РРРР-ММ-ДД]
# Результат відновлюється в Каса → Налаштування → «Відновити з файлу».
import openpyxl, sys, re, json, collections, datetime, random, string

SRC, OUT, REPORT_DATE = sys.argv[1], sys.argv[2], (sys.argv[3] if len(sys.argv) > 3 else "")
wb = openpyxl.load_workbook(SRC)
norm = lambda s: re.sub(r"\s+", " ", str(s or "")).strip()
key = lambda s: norm(s).lower()
_n = 0
def uid():
    global _n; _n += 1
    return "x" + format(_n, "05d") + "".join(random.choices(string.ascii_lowercase + string.digits, k=5))
iso = lambda dt: dt.strftime("%Y-%m-%dT%H:%M:%S.000Z")
NOW = datetime.datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%S.000Z")

# ---- Учасники ----
ref = [norm(r[0].value) for r in wb["Довідники"].iter_rows(min_row=2) if r[0].value]
NOT_PEOPLE = {"волонтери"}
ALIAS = {"гайша андрій геннадійович": "гайша андрій геннадійович ворон"}
people, pid = [], {}
for n in ref:
    if key(n) in NOT_PEOPLE: continue
    p = {"id": uid(), "name": n, "phone": "", "note": "", "active": True, "createdAt": NOW}
    people.append(p); pid[key(n)] = p["id"]

# ---- Журнал ----
rows = []
for r in wb["Журнал"].iter_rows(min_row=2):
    v = [c.value for c in r]
    if not v[0]: continue
    rows.append((r[0].row, v))
rows.sort(key=lambda x: (x[1][0], x[0]))

def amount(v):
    if isinstance(v, (int, float)): return round(float(v), 2)
    if isinstance(v, str) and v.startswith("="):  # напр. =7000+17500
        return round(float(eval(v[1:], {"__builtins__": {}})), 2)
    return round(float(str(v).replace(",", ".")), 2)

CATS = [
    ("Старлінк", r"старлінк|starlink|старлинк"),
    ("Підписки та сервіси", r"інтернет"),
    ("Поповнення банок майстерень", r"поповнення|на майстерню бпла|майстерн\w* (бпла|нрк)|^майстерня"),
    ("Розрахунки з учасниками", r"^(борг |від \w+ до|\w+ від \w+$)|барні"),
    ("Оренда", r"оренд"),
    ("Підписки та сервіси", r"підписк|промоплат|пром |аліекспрес"),
    ("Авто", r"щеплен|балансир|кардан|рвд|омивач|олива|lt ?46|ml ?270|тяга|гальм|карбюр|\bт4\b|\bт5\b|л200|навар|хайлюкс|hilux|хайс|рено|\bjac\b|таракан|шин|колес|розвал|запчаст|масло|фильтр|фільтр|колодк|стаб|автомоб|маховик|присадк|ремінь|ремень|\bгур\b|втулк|шиномонтаж|тонув|акпп|мкпп|двигун|аккумулятор авто"),
    ("Живлення та паливо", r"заправка (?!балон)|\bдп\b|біопалив|генератор|бензин|палив|пальне|зарядн|ecoflow|ecoplay|акум|батаре|інвертор"),
    ("БПЛА та зв'язок", r"антен|модуль|мавік|mavic|матріс|matrice|дрон|бпла|пропел|пропол|вампір|пульт|crossfire|hdmi|sll|бустер|спарк|шлейф|роутер|mesh|рації|рація"),
    ("Техніка та електроніка", r"нрк|р2д2|діод|медіаконвер|патчкорд|ноутбук|ssd|пам.ять|плат[аи]|комплектуюч|пластик abs|abs|фпв|пропи|мікроскоп|мультиметр|паяль|адаптер|usb|сімкарт|принтер|картридж|кадрідж|запобіжник|відеозахват|bluetooth"),
    ("Доставка", r"нова пошта|\bнп\b|доставк|послуг"),
    ("Господарські", r"стяжк|скоч|скотч|розетк|вилк|кабель|кабеля|лампоч|рукавиц|стрічк|клей|пакет|дюбел|інструмент|ключ|бур|мішк|фарб|пінофол|подовжувач|піна|труб|будматер|скоби|цвях|брус|перчатк|пупирк|канцеляр|стіл|стільч|тележк|дистил|прапор|омивач"),
]
def category(c):
    s = key(c)
    for name, rx in CATS:
        if re.search(rx, s): return name
    return "Інше"

tx, unknown = [], collections.Counter()
for rownum, v in rows:
    ts, typ = v[0], v[1]
    base = {"id": uid(), "date": ts.date().isoformat(), "createdAt": iso(ts), "srcRow": rownum}
    if typ == "Прихід":
        name = key(v[2]); name = ALIAS.get(name, name)
        comment = norm(v[4])
        t = {**base, "type": "income", "amount": amount(v[3]), "personId": "", "collectionId": "", "category": "", "comment": comment}
        if name in NOT_PEOPLE:
            t["comment"] = norm(v[2]) + (": " + comment if comment else "")
        else:
            if name not in pid:  # є в журналі, але немає в довіднику — колишні учасники
                p = {"id": uid(), "name": norm(v[2]), "phone": "", "note": "немає в довіднику", "active": False, "createdAt": NOW}
                people.append(p); pid[name] = p["id"]; unknown[p["name"]] += 0
            t["personId"] = pid[name]
        tx.append(t)
    elif typ == "Витрата":
        c = norm(v[6])
        tx.append({**base, "type": "expense", "amount": amount(v[5]), "personId": "", "collectionId": "", "category": category(c), "comment": c})

# ---- Збори (хвилі внесків) ----
contrib = [t for t in tx if t["type"] == "income" and t["personId"] and not t["comment"]]
perday = collections.Counter(t["date"] for t in contrib)
waves = []
for d in sorted(perday):
    if perday[d] < 20: continue
    dd = datetime.date.fromisoformat(d)
    if waves and (dd - datetime.date.fromisoformat(waves[-1]["days"][-1])).days <= 2:
        waves[-1]["days"].append(d); continue
    waves.append({"days": [d]})

for w in waves:
    amts = [t["amount"] for t in contrib if t["date"] in w["days"]]
    share2000 = sum(a == 2000 for a in amts) / len(amts)
    if w is waves[0]: w["kind"] = "other"
    else: w["kind"] = "salary" if share2000 >= 0.7 else "bonus"
    w["date"] = w["days"][0]

def wave_for(t):
    cur = None
    for w in waves:
        if w["date"] <= t["date"]: cur = w
    return cur

paid = collections.defaultdict(set)
for t in sorted(contrib, key=lambda t: t["createdAt"]):
    w = wave_for(t)
    if not w: continue
    i = waves.index(w)
    # Запізнілий внесок за хвилю іншого типу (напр. 2000 за ЗП у період премії)
    want_salary = t["amount"] == 2000
    if w["kind"] != "other" and (w["kind"] == "salary") != want_salary and t["date"] not in w["days"]:
        prev = next((x for x in reversed(waves[:i]) if x["kind"] == ("salary" if want_salary else "bonus")), None)
        if prev and t["personId"] not in paid[id(prev)]: w = prev
    t["collectionId"] = w.setdefault("id", uid())
    paid[id(w)].add(t["personId"])

KL = {"salary": "ЗП", "bonus": "премію", "other": "внесок"}
collections_out = []
active = {p["id"] for p in people if p["active"]}
for i, w in enumerate(waves):
    w.setdefault("id", uid())
    y, m, d = w["date"].split("-")
    title = f"Внесок за {KL[w['kind']]} {d}.{m}.{y}" if w["kind"] != "other" else f"Перший внесок {d}.{m}.{y}"
    ids = set(paid[id(w)])
    is_recent = i >= len(waves) - 2
    if is_recent:  # для актуальних зборів — очікувані = хто платив у попередньому зборі того ж типу
        prev = next((x for x in reversed(waves[:i]) if x["kind"] == w["kind"]), None)
        if prev: ids |= (paid[id(prev)] & active)
    order = {p["id"]: n for n, p in enumerate(people)}
    collections_out.append({
        "id": w["id"], "title": title, "kind": w["kind"], "date": w["date"],
        "amount": 2000 if w["kind"] == "salary" else 0,
        "personIds": sorted(ids, key=lambda x: order[x]), "comment": "імпорт з таблиці",
        "closed": not is_recent, "createdAt": NOW,
    })

# ---- Гаманці підрозділів (окремі аркуші-банки) ----
WALLETS = [("Банка Стрікс", "майстерня НРК", r"нрк"), ("Банка Крістал", "майстерня БПЛА", r"бпла")]
wallets, wtx_all = [], []
for name, note, hint in WALLETS:
    if name not in wb.sheetnames: continue
    w = {"id": uid(), "name": name, "note": note, "openingBalance": 0, "active": True, "createdAt": NOW, "hint": hint}
    wallets.append(w)
    raw = [(r[0].row, [c.value for c in r]) for r in wb[name].iter_rows(min_row=3)]
    raw = [(n, v) for n, v in raw if any(x not in (None, "", " ") for x in v[:4])]
    first = next(v[0] for _, v in raw if isinstance(v[0], datetime.datetime))
    w["lastDate"] = max(v[0] for _, v in raw if isinstance(v[0], datetime.datetime)).date().isoformat()
    last_dt = first
    for n, v in raw:
        dt = v[0] if isinstance(v[0], datetime.datetime) else last_dt  # рядок без дати — дата попереднього
        last_dt = dt
        c = norm(v[3])
        ts = iso(dt.replace(hour=12) + datetime.timedelta(seconds=n))
        base = {"date": dt.date().isoformat(), "walletId": w["id"], "personId": "", "collectionId": "", "createdAt": ts}
        inc_v, exp_v = v[1], v[2]
        if isinstance(inc_v, (int, float)) and inc_v:
            # якщо в рядку є і прихід, і витрата — коментар стосується витрати
            tc = "" if isinstance(exp_v, (int, float)) and exp_v else c
            if not isinstance(v[0], datetime.datetime): tc = (tc + " " if tc else "") + "(у таблиці без дати)"
            wtx_all.append({**base, "id": uid(), "type": "income", "amount": amount(inc_v), "category": "", "comment": tc})
        if isinstance(exp_v, (int, float)) and exp_v:
            wtx_all.append({**base, "id": uid(), "type": "expense", "amount": amount(exp_v), "category": category(c), "comment": c})

# Зіставлення: поповнення в касі ↔ надходження в банці → зв'язаний переказ
d = lambda x: datetime.date.fromisoformat(x)
main_top = [t for t in tx if t["type"] == "expense" and t["category"] == "Поповнення банок майстерень"]
used, matched = set(), 0
w_by_id = {w["id"]: w for w in wallets}
for wi in sorted([t for t in wtx_all if t["type"] == "income"], key=lambda t: t["date"]):
    refundish = wi["comment"] and "поповнен" not in wi["comment"].lower()  # повернення, компенсації
    hint = w_by_id[wi["walletId"]]["hint"]
    best = None
    for m in main_top:
        if m["id"] in used or m["amount"] != wi["amount"]: continue
        gap = abs((d(m["date"]) - d(wi["date"])).days)
        own = bool(re.search(hint, key(m["comment"])))
        if gap > (20 if own else 10) or (refundish and not own): continue
        other = any(re.search(x["hint"], key(m["comment"])) for x in wallets if x is not w_by_id[wi["walletId"]])
        if other and not own: continue
        score = (0 if own else 1, gap)
        if not best or score < best[0]: best = (score, m)
    if best:
        m = best[1]; used.add(m["id"]); matched += 1
        tid = uid()
        m.update({"transferId": tid, "peerWallet": wi["walletId"], "category": "Переказ"})
        wi.update({"transferId": tid, "peerWallet": "", "comment": m["comment"] or wi["comment"]})
# Поповнення, записані в касі після останнього рядка таблиці банки, — теж перекази
late = 0
for m in main_top:
    if m["id"] in used: continue
    w = next((x for x in wallets if re.search(x["hint"], key(m["comment"]))), None)
    if not w or m["date"] <= w["lastDate"]: continue
    tid = uid(); used.add(m["id"]); late += 1
    m.update({"transferId": tid, "peerWallet": w["id"], "category": "Переказ"})
    wtx_all.append({"id": uid(), "type": "income", "amount": m["amount"], "date": m["date"], "walletId": w["id"],
                    "personId": "", "collectionId": "", "category": "", "createdAt": m["createdAt"], "transferId": tid, "peerWallet": "",
                    "comment": m["comment"] + " (немає в таблиці банки — з журналу каси)"})
unmatched_main = [t for t in main_top if t["id"] not in used]
unmatched_w = [t for t in wtx_all if t["type"] == "income" and not t.get("transferId")]
w_names = {w["id"]: w["name"] for w in wallets}
for w in wallets: w.pop("hint"); w.pop("lastDate")
tx.extend(wtx_all)

# ---- Звіт-відсічка ----
reports = []
if REPORT_DATE:
    items = [t for t in tx if t["date"] <= REPORT_DATE and not t.get("walletId")]
    inc = round(sum(t["amount"] for t in items if t["type"] == "income"), 2)
    exp = round(sum(t["amount"] for t in items if t["type"] == "expense"), 2)
    r = {"id": uid(), "walletId": "", "number": 1, "date": REPORT_DATE, "prevDate": "", "opening": 0, "income": inc, "expense": exp,
         "closing": round(inc - exp, 2), "txIds": [t["id"] for t in items], "comment": "Підсумок старої таблиці (імпорт)", "createdAt": NOW}
    for t in items: t["reportId"] = r["id"]
    reports.append(r)

for t in tx: t.pop("srcRow", None)
cats = [c for c, _ in CATS] + ["Інше"]
for t in tx: t.setdefault("walletId", "")
db = {"people": people, "tx": tx, "collections": collections_out, "reports": reports, "wallets": wallets,
      "settings": {"title": "Каса", "monthlyFee": 0, "openingBalance": 0, "categories": cats}}
json.dump(db, open(OUT, "w"), ensure_ascii=False)

# ---- Підсумок ----
for w in [{"id": "", "name": "Каса (основна)"}] + wallets:
    lst = [t for t in tx if t["walletId"] == w["id"]]
    i_, e_ = sum(t["amount"] for t in lst if t["type"] == "income"), sum(t["amount"] for t in lst if t["type"] == "expense")
    print(f"{w['name']:<18} операцій {len(lst):>5}  прихід {i_:>12.2f}  витрати {e_:>12.2f}  залишок {i_-e_:>10.2f}")
print(f"переказів зіставлено: {matched}, додано пізніших за таблицю банки: {late}")
print("поповнення в касі без пари в банці:"); [print("   ", t["date"], t["amount"], t["comment"]) for t in unmatched_main]
print("надходження в банках без пари в касі:"); [print("   ", w_names[t["walletId"]], t["date"], t["amount"], t["comment"]) for t in unmatched_w]
tx = [t for t in tx if not t["walletId"]]
inc = sum(t["amount"] for t in tx if t["type"] == "income"); exp = sum(t["amount"] for t in tx if t["type"] == "expense")
print(f"учасників {len(people)} (активних {len(active)}, з журналу без довідника: {len(unknown)}: {', '.join(unknown)})")
print(f"операцій {len(tx)}: прихід {inc:.2f} ({sum(t['type']=='income' for t in tx)}), витрати {exp:.2f} ({sum(t['type']=='expense' for t in tx)}), залишок {inc-exp:.2f}")
print("без учасника (волонтери тощо):", sum(1 for t in tx if t['type']=='income' and not t['personId']),
      "| внески з коментарем (повернення тощо, без збору):", sum(1 for t in tx if t['type']=='income' and t['personId'] and t['comment']))
for c in collections_out:
    s = sum(t["amount"] for t in tx if t["collectionId"] == c["id"]); n = len({t["personId"] for t in tx if t["collectionId"] == c["id"]})
    print(f"  {c['title']:<34} внесли {n:>3} з {len(c['personIds']):>3}  {s:>10.0f}  {'закр' if c['closed'] else 'ВІДКРИТИЙ'}")
print("категорії:", collections.Counter(t["category"] for t in tx if t["type"] == "expense").most_common())
for r in reports: print("звіт", r["date"], "прихід", r["income"], "витрати", r["expense"], "залишок", r["closing"], "операцій", len(r["txIds"]))
